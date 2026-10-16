#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""策划数据 XLSX → 分系统 CSV → 已接入的 Godot Resource。标准库 + openpyxl。"""
from __future__ import annotations

import argparse
import csv
import io
import json
import math
import re
import sys
from collections import defaultdict
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
SHEET_NAMES = (
    "01_身份配置", "02_主播关卡", "03_普通词库", "04_矛盾内容", "05_关卡生成",
    "06_战斗数值", "06_生命周期", "06_基础参数", "07_特性参数", "08_Tier档位", "09_直播数据", "10_复读配置",
    "11_矛盾参数", "12_神谕参数", "13_吞并奖励", "14_败者卡", "15_神降临",
    "16_结局配置", "17_音频事件", "18_表现资产",
    "19_直播评论词库", "20_主播气泡对白", "21_反击技能池",
    "22_直播评论表现参数", "23_对手立绘配置", "24_新手教学配置",
    "25_直播观众账号",
)
ADDITIONAL_SHEETS = frozenset(SHEET_NAMES[-7:])
# 原字段名是唯一输入契约；这里只约束当前已确认的主键、核心字段、类型和合法枚举。
RULES = {
    "01_身份配置": ("identity_id", ("identity_id", "display_name", "tendency")),
    "02_主播关卡": ("level_id", ("level_id", "level_order", "streamer_id")),
    "03_普通词库": ("word_id", ("word_id", "text", "tendency", "strength", "weight", "pool_id", "enabled")),
    "04_矛盾内容": ("contradiction_id", ("contradiction_id", "set_id", "streamer_id", "type", "text")),
    "05_关卡生成": ("level_id", ("level_id",)),
    "06_战斗数值": ("param_id", ("param_id", "value")),
    "06_生命周期": ("barrage_type", ("barrage_type", "base_lifetime_s")),
    "06_基础参数": ("param_id", ("param_id", "value")),
    "07_特性参数": ("trait_id", ("trait_id",)),
    "08_Tier档位": ("tier", ("tier", "pk_up_threshold", "spawn_batch_mult", "spawn_frequency_mult", "move_speed_mult", "lifetime_mult", "pullback_mult", "repeat_count", "repeat_lifetime_s")),
    "09_直播数据": ("rule_id", ("rule_id",)),
    "10_复读配置": ("param_id", ("param_id",)),
    "11_矛盾参数": ("level_id", ("level_id",)),
    "12_神谕参数": ("param_id", ("param_id",)),
    "13_吞并奖励": ("reward_id", ("reward_id",)),
    "14_败者卡": ("card_id", ("card_id",)),
    "15_神降临": ("param_id", ("param_id",)),
    "16_结局配置": (None, ("config_type", "primary_tendency")),
    "17_音频事件": ("event_id", ("event_id",)),
    "18_表现资产": ("asset_id", ("asset_id",)),
    "19_直播评论词库": ("comment_id", ("comment_id", "viewer_id", "text", "enabled")),
    "20_主播气泡对白": ("dialogue_id", ("dialogue_id", "streamer_id", "speaker_side", "trigger_type", "text", "enabled")),
    "21_反击技能池": ("skill_id", ("skill_id", "streamer_id", "tier", "trait_ids", "enabled")),
    "22_直播评论表现参数": ("param_id", ("param_id", "value")),
    "23_对手立绘配置": ("portrait_set_id", ("portrait_set_id", "streamer_id", "idle_asset_path", "enabled")),
    "24_新手教学配置": ("step_id", ("tutorial_config_id", "step_id", "step_order", "enabled")),
    "25_直播观众账号": ("viewer_id", ("viewer_id", "display_name", "audience_type", "enabled")),
}
# 显式数值/布尔类型边界，避免将 ID 中的前导零、文本中的标点意外转换。
NUMERIC = {
    "02_主播关卡": {"level_order": "int"},
    "03_普通词库": {"strength": "int", "weight": "float"},
    "05_关卡生成": {"base_batch_count": "int", "spawn_interval_s": "float", "move_speed_px_s": "float", "word_weight_multiplier": "float"},
    "06_生命周期": {"base_lifetime_s": "float"},
    "07_特性参数": {"lifetime_override_s": "float", "split_child_count": "int"},
    "08_Tier档位": {"tier": "int", "pk_up_threshold": "float", "pk_down_threshold": "float", "spawn_batch_mult": "float", "spawn_frequency_mult": "float", "move_speed_mult": "float", "lifetime_mult": "float", "pullback_mult": "float", "repeat_count": "int", "repeat_lifetime_s": "float"},
    "09_直播数据": {"viewer_value": "float", "like_value": "float", "comment_delta": "float", "fan_value": "float"},
    "11_矛盾参数": {"time_limit_s": "float", "shot_limit": "int", "silence_transition_s": "float"},
    "13_吞并奖励": {"inherited_weight": "float"},
    "17_音频事件": {"volume_db": "float", "pitch_min": "float", "pitch_max": "float"},
    "19_直播评论词库": {"weight": "float"},
    "20_主播气泡对白": {"line_order": "int", "display_duration_s": "float", "trigger_time_s": "float", "random_weight": "float"},
    "21_反击技能池": {"tier": "int", "weight": "float"},
    "24_新手教学配置": {"step_order": "int"},
    "25_直播观众账号": {"weight": "float"},
}
BOOL_FIELDS = {"enabled", "inheritable", "loop", "affected_by_tier_lifetime_mult"}
BOOL_PARAM_IDS = {"comment_pause_with_game", "comment_updates_business_counter"}
ENUMS = {
    ("01_身份配置", "tendency"): {"orthodox", "heretical", "heresy", "absurd"},
    ("03_普通词库", "tendency"): {"orthodox", "heretical", "heresy", "absurd", "neutral"},
    ("04_矛盾内容", "type"): {"true", "false"},
}
# 进入现有 Resource 时唯一允许的字段重命名，游戏层不用猜策划的列名。
SPEECH_TO_GODOT = {
    "word_id": "original_sentence_id", "text": "text", "tendency": "tendency_id",
    "strength": "strength", "weight": "appearance_weight",
}
TIER_TO_GODOT = {
    "pk_up_threshold": "upgrade_threshold",
    "pk_down_threshold": "downgrade_threshold",
    "spawn_batch_mult": "generation_count_multiplier",
    "spawn_frequency_mult": "generation_frequency_multiplier",
    "move_speed_mult": "movement_speed_multiplier",
    "lifetime_mult": "lifetime_multiplier",
    "pullback_mult": "opponent_pullback_multiplier",
    "repeat_count": "repeat_count_per_hit",
    "motion_straight_weight": "straight_motion_weight",
    "motion_curve_weight": "curved_motion_weight",
    "motion_speed_change_weight": "speed_change_motion_weight",
    "motion_wander_weight": "wandering_motion_weight",
}
MOTION_WEIGHT_FIELDS = (
    "motion_straight_weight", "motion_curve_weight",
    "motion_speed_change_weight", "motion_wander_weight",
)
BASIC_TO_GODOT = {"spawn_interval_s": "base_spawn_interval_seconds"}
PK_PERCENT_POINT_SCALE = 0.01  # 策划 0.12 个百分点 → Godot 0.0012。


def value_text(value: object) -> str:
    """保留文本原貌；仅将 Excel 原生数字和布尔值转换为稳定可读的 CSV 字符串。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return format(value, ".15g") if math.isfinite(value) else str(value)
    return str(value)


def is_header(row: list[str], required: tuple[str, ...]) -> bool:
    """按照真实字段名匹配表头，避免误把区块标题当作记录。"""
    filled = [s.strip() for s in row if s.strip()]
    return bool(filled) and len(filled) == len(set(filled)) and all(f in filled for f in required)


def xlsx_rows(ws) -> list[tuple[int, list[str]]]:
    """读取物理行号；对空行保留行号，保证错误可追溯到 Excel。"""
    result = []
    for n, cells in enumerate(ws.iter_rows(values_only=True), 1):
        row = [value_text(cell) for cell in cells]
        while row and not row[-1]:
            row.pop()
        result.append((n, row))
    return result


def bool_text(text: str) -> str | None:
    """只接受明确的布尔文本；不把未知值默认为真。"""
    return {"true": "true", "false": "false", "1": "true", "0": "false", "yes": "true", "no": "false"}.get(text.strip().lower())


def validate_native_boolean_cells(ws, report: Report) -> None:
    """新版正式 Google 表必须保留 Excel 原生布尔类型，避免迁移把真假变成字符串。

    旧 TEST_ONLY/历史工作簿无22表，保持旧格式的兼容导出。
    """
    iterator = ws.iter_rows(values_only=True)
    header = next(iterator, ())
    if not header:
        return
    mapping = {str(field): i for i, field in enumerate(header) if field is not None}
    bool_columns = [(field, idx) for field, idx in mapping.items() if field in BOOL_FIELDS]
    value_col = mapping.get("value") if ws.title == "22_直播评论表现参数" else None
    key_col = mapping.get("param_id", 0)
    for row_num, row in enumerate(iterator, 2):
        if not row or not str(row[0] or "").strip():
            continue
        if str(row[0]).startswith(("填写", "【", "EXAMPLE_")):
            continue
        for field, idx in bool_columns:
            value = row[idx] if idx < len(row) else None
            if value is not None and type(value) is not bool:
                report.add("错误", ws.title, row_num, field, "必须是原生布尔单元格（复选框TRUE/FALSE），不能是文本")
        if value_col is None:
            continue
        key = str(row[key_col] if key_col < len(row) else "")
        value = row[value_col] if value_col < len(row) else None
        if value is None:
            continue
        if key in BOOL_PARAM_IDS:
            if type(value) is not bool:
                report.add("错误", ws.title, row_num, "value", f"{key} 必须是原生布尔开关")
        elif type(value) is bool or type(value) not in (int, float) or not math.isfinite(value):
            report.add("错误", ws.title, row_num, "value", f"{key} 必须是数值参数")


class Report:
    """汇总精确位置及分类；警告允许导出，错误会终止本轮写入。"""
    def __init__(self):
        self.errors: list[str] = []
        self.warnings: list[str] = []
        self.stats: list[tuple[str, int, int, int]] = []

    def add(self, level: str, sheet: str, row: int, field: str, why: str) -> None:
        line = f"{sheet} 第{row}行 [{field}] {why}"
        (self.errors if level == "错误" else self.warnings).append(line)


def parse_sheet(sheet: str, rows: list[tuple[int, list[str]]], report: Report, start: int = 0, end: int | None = None):
    """提取一块字段结构固定的表；只过滤明确的示例/说明/全空行。"""
    key, required = RULES[sheet]
    block = rows[start:end]
    header_idx = next((i for i, (_, row) in enumerate(block) if is_header(row, required)), None)
    if header_idx is None:
        report.add("错误", sheet, 1, "表头", "缺少正式表头或必填列：" + ", ".join(required))
        return [], []
    header_row_no, header = block[header_idx]
    if len(header) != len(set(header)) or any(not f.strip() for f in header):
        report.add("错误", sheet, header_row_no, "表头", "字段名为空或重复")
        return [], []
    records: list[tuple[int, dict[str, str]]] = []
    blanks = examples = 0
    for line, row in block[header_idx + 1:]:
        if not row or not any(s.strip() for s in row):
            blanks += 1
            continue
        if row[0].startswith("【") or is_header(row, required):
            report.add("警告", sheet, line, "区块", "跳过重复标题或表头")
            continue
        # 02 表底部有逐列的“填写……”说明，按照说明型文案而非位置识别。
        guide_count = sum(s.strip().startswith(("填写", "给每", "根据", "仅在", "选择")) for s in row)
        if row[0].strip().startswith(("填写", "给每个")) and guide_count >= 1:
            examples += 1
            report.add("警告", sheet, line, "说明行", "已跳过字段填写说明")
            continue
        record = {field: row[i] if i < len(row) else "" for i, field in enumerate(header)}
        primary = record.get(key, "") if key else ""
        if primary.upper().startswith("EXAMPLE_") or record.get("notes", "").strip().startswith("示例行"):
            examples += 1
            report.add("警告", sheet, line, key or "notes", "已跳过明确标记的示例行")
            continue
        if len(row) > len(header) and any(s.strip() for s in row[len(header):]):
            report.add("错误", sheet, line, "列数", "记录超出了字段表头长度")
        records.append((line, record))
    report.stats.append((sheet, len(records), blanks, examples))
    return header, records


def numeric(text: str, kind: str) -> bool:
    """验算数值合法性：整数不接受 1.5；浮点必须有限。"""
    try:
        number = float(text)
        if not math.isfinite(number):
            return False
        return kind != "int" or number.is_integer()
    except ValueError:
        return False


def validate(tables: dict, report: Report) -> None:
    """集中检查必填、类型、枚举、稳定 ID 和已具备双方来源的关联。"""
    for name, (header, records) in tables.items():
        key, required = RULES[name]
        seen: set[str] = set()
        for line, record in records:
            for f in required:
                if record.get(f, "").strip():
                    continue
                # 允许未完成参数以警告形式保留空值；稳定 ID、台词内容仍为硬错误。
                severity = "错误" if f == key or (name == "03_普通词库" and f in ("text", "tendency", "strength", "weight", "pool_id", "enabled")) else "警告"
                report.add(severity, name, line, f, "必填内容尚未填写")
            if key:
                identity = record.get(key, "").strip()
                if identity:
                    if identity in seen:
                        report.add("错误", name, line, key, f"稳定 ID 重复：{identity}")
                    seen.add(identity)
            for f, kind in NUMERIC.get(name, {}).items():
                raw = record.get(f, "").strip()
                if raw and not numeric(raw, kind):
                    report.add("错误", name, line, f, f"需要{kind}类型，实际为 {raw!r}")
            for f in header:
                raw = record.get(f, "").strip()
                if f in BOOL_FIELDS and raw:
                    parsed = bool_text(raw)
                    if parsed is None:
                        report.add("错误", name, line, f, f"非法布尔值：{raw!r}")
                    else:
                        record[f] = parsed
                allowed = ENUMS.get((name, f))
                if allowed and raw and raw not in allowed:
                    report.add("错误", name, line, f, f"枚举值 {raw!r} 不在 {sorted(allowed)}")
                if f not in ("text", "notes", "description", "clue_text") and raw.startswith(("待填写", "填写", "TODO", "TBD")):
                    report.add("警告", name, line, f, "策划占位，尚待填写")
            # 只报告已确认应填写的配置空值；未完工内容保留空白，绝不造数。
            pending_fields = {
                "05_关卡生成": ("base_batch_count", "spawn_interval_s", "move_speed_px_s", "normal_screen_cap"),
                "10_复读配置": ("value",),
                "15_神降临": ("value",),
                "17_音频事件": ("asset_path_or_id",),
            }.get(name, ())
            for f in pending_fields:
                if f in header and not record.get(f, "").strip():
                    report.add("警告", name, line, f, "正式值尚待策划或资产交付")
            if name == "22_直播评论表现参数":
                # 数据列虽共用value，类型由param_id决定，而非任意允许混填。
                raw = record.get("value", "")
                if record.get("param_id", "") in BOOL_PARAM_IDS:
                    if raw and raw.lower() not in ("true", "false"):
                        report.add("错误", name, line, "value", "布尔参数只能为 true/false")
                elif raw and not numeric(raw, "float"):
                    report.add("错误", name, line, "value", "数值参数必须是有限数字")
            if name == "03_普通词库":
                strength = record.get("strength", "")
                weight = record.get("weight", "")
                if numeric(strength, "int") and int(float(strength)) not in (1, 2, 3):
                    report.add("错误", name, line, "strength", "强度只能为 1、2、3")
                if numeric(weight, "float") and float(weight) <= 0:
                    report.add("错误", name, line, "weight", "出现权重必须大于 0")
    # 新增表仅在实际提供时校验关联；旧TEST_ONLY夹具保持兼容。
    for child, field, parent, target in (
        ("02_主播关卡", "portrait_set_id", "23_对手立绘配置", "portrait_set_id"),
        ("02_主播关卡", "tutorial_config_id", "24_新手教学配置", "tutorial_config_id"),
        ("02_主播关卡", "story_config_id", "20_主播气泡对白", "story_config_id"),
        ("19_直播评论词库", "viewer_id", "25_直播观众账号", "viewer_id"),
        ("20_主播气泡对白", "streamer_id", "02_主播关卡", "streamer_id"),
        ("21_反击技能池", "streamer_id", "02_主播关卡", "streamer_id"),
        ("23_对手立绘配置", "streamer_id", "02_主播关卡", "streamer_id"),
        ("05_关卡生成", "level_id", "02_主播关卡", "level_id"),
        ("11_矛盾参数", "level_id", "02_主播关卡", "level_id"),
        ("04_矛盾内容", "streamer_id", "02_主播关卡", "streamer_id"),
        ("07_特性参数", "penalty_param_id", "06_战斗数值", "param_id"),
        ("02_主播关卡", "word_pool_id", "03_普通词库", "pool_id"),
        ("13_吞并奖励", "streamer_id", "02_主播关卡", "streamer_id"),
    ):
        if child not in tables or parent not in tables:
            continue
        source = {r.get(target, "") for _, r in tables[parent][1] if r.get(target)}
        if not source:
            for line, row in tables[child][1]:
                if row.get(field, "").strip():
                    report.add("警告", child, line, field, f"{parent} 目前没有正式记录，关联 ID 暂无法核验")
            continue
        for line, row in tables[child][1]:
            # 停用的吞并奖励占位记录保留在 CSV，正式关联仅核验启用项。
            if (child in ("13_吞并奖励", "19_直播评论词库",
                          "20_主播气泡对白", "21_反击技能池", "23_对手立绘配置")
                    and row.get("enabled", "").lower() == "false"):
                continue
            ref = row.get(field, "").strip()
            if ref and ref not in source:
                report.add("错误", child, line, field, f"关联 ID {ref!r} 未出现在 {parent}.{target}")
    # 表格的 Tier 数值单位是 PK 百分点，Godot 阈值使用 [0,1]；这里同时验算趋势。
    if "08_Tier档位" not in tables:
        return
    ts = tables["08_Tier档位"][1]
    if len(ts) != 6 or {r["tier"] for _, r in ts} != set("012345"):
        report.add("错误", "08_Tier档位", 1, "tier", "需要完整 Tier 0～5 各一条")
    previous = -1.0
    for line, r in ts:
        up, down = r["pk_up_threshold"], r.get("pk_down_threshold", "")
        if not numeric(up, "float"):
            continue
        upper = float(up)
        if not (0 < upper <= 100 and upper > previous):
            report.add("错误", "08_Tier档位", line, "pk_up_threshold", "升档阈值需递增，范围 (0,100]")
        previous = upper
        if r["tier"] != "0" and (not down or not numeric(down, "float")):
            report.add("错误", "08_Tier档位", line, "pk_down_threshold", "Tier 1～5 必须填写降档阈值")
        if down and numeric(down, "float") and not (0 <= float(down) < upper):
            report.add("错误", "08_Tier档位", line, "pk_down_threshold", "降档阈值须低于本档升档阈值")
        motion_values = [r.get(f, "") for f in MOTION_WEIGHT_FIELDS]
        if any(motion_values):
            if not all(numeric(x, "int") for x in motion_values):
                report.add("错误", "08_Tier档位", line, "motion", "四种运动权重须全部为整数")
            else:
                weights = [int(float(x)) for x in motion_values]
                if min(weights) < 0 or sum(weights) != 100:
                    report.add("错误", "08_Tier档位", line, "motion", "四项非负运动权重之和必须等于100")
        for f in ("spawn_batch_mult", "spawn_frequency_mult", "move_speed_mult", "lifetime_mult", "pullback_mult"):
            if not numeric(r.get(f, ""), "float"):
                continue
            value = float(r[f])
            # T0 普通战斗无回拉；从 T1 起回拉倍率必须为正。
            if f == "pullback_mult" and r["tier"] == "0":
                if value < 0:
                    report.add("错误", "08_Tier档位", line, f, "T0 回拉倍率不得小于 0")
            elif value <= 0:
                report.add("错误", "08_Tier档位", line, f, "倍率必须大于 0")
        if numeric(r["repeat_count"], "int") and int(float(r["repeat_count"])) < 0:
            report.add("错误", "08_Tier档位", line, "repeat_count", "复读数量不可为负")
        if numeric(r["repeat_lifetime_s"], "float") and numeric(r["lifetime_mult"], "float"):
            expected = 3.0 * float(r["lifetime_mult"])
            if abs(float(r["repeat_lifetime_s"]) - expected) > 0.001:
                report.add("错误", "08_Tier档位", line, "repeat_lifetime_s", f"应等于 3s×lifetime_mult={expected:g}")


def csv_body(header: list[str], records: list) -> str:
    """CSV 使用标准引号转义、UTF-8、固定换行；重复导出字节稳定。"""
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(header)
    for _, record in records:
        writer.writerow([record.get(f, "") for f in header])
    return output.getvalue()


def gd_string(s: str) -> str:
    """Godot 资源文本以 JSON 规则转义引号、反斜杠和换行。"""
    return json.dumps(s, ensure_ascii=False)


def make_speech_pools(records: list, report: Report, output_prefix: str) -> dict[str, str]:
    """按 pool_id 生成 LevelSpeech 的真实 Resource 子资源；停用行仍留在 CSV。"""
    grouped: dict[str, list] = defaultdict(list)
    for line, r in records:
        if r.get("enabled", "").lower() != "true":
            continue
        pool = r["pool_id"]
        if not re.fullmatch(r"[A-Za-z0-9_-]+", pool):
            report.add("错误", "03_普通词库", line, "pool_id", "Resource 文件名只允许英文、数字、_ 和 -")
            continue
        grouped[pool].append(r)
    results = {}
    for pool, rows in grouped.items():
        parts = [
            f'[gd_resource type="Resource" script_class="LevelSpeechPool" load_steps={len(rows)+3} format=3]',
            "",
            '[ext_resource type="Script" path="res://data/level_configuration/level_speech_pool.gd" id="1_pool"]',
            '[ext_resource type="Script" path="res://data/level_configuration/level_speech.gd" id="2_speech"]',
            "",
        ]
        ref_ids = []
        for i, r in enumerate(rows, 1):
            rid = f"Speech_{i}"
            ref_ids.append(f'SubResource("{rid}")')
            parts.extend([
                f'[sub_resource type="Resource" script_class="LevelSpeech" id="{rid}"]',
                'script = ExtResource("2_speech")',
                f'original_sentence_id = {gd_string(r["word_id"])}',
                f'text = {gd_string(r["text"])}',
                f'tendency_id = {gd_string("heretical" if r["tendency"] == "heresy" else r["tendency"])}',
                f'strength = {int(float(r["strength"]))}',
                f'appearance_weight = {format(float(r["weight"]), ".15g")}',
                "",
            ])
        parts.extend([
            "[resource]", 'script = ExtResource("1_pool")',
            f"pool_id = {gd_string(pool)}",
            'speeches = Array[ExtResource("2_speech")]([' + ", ".join(ref_ids) + "])",
            "",
        ])
        results[f"{output_prefix}/{pool}.tres"] = "\n".join(parts)
    return results


def make_tier_resource(records: list, report: Report) -> str:
    """用已有正式 .tres 作为模板，只替换 Sheet 已覆盖的字段，保留 Neutral 与立绘状态。"""
    template_path = ROOT / "data/combat_stage/tier_catalog.tres"
    source = template_path.read_text(encoding="utf-8")
    for line, row in records:
        tier = row["tier"]
        if not numeric(tier, "int"):
            continue
        tier = str(int(float(tier)))
        pattern = re.compile(r'(\[sub_resource type="Resource" id="Tier' + tier + r'"\]\r?\n)(.*?)(?=\r?\n\[(?:sub_resource|resource)(?:\s|\])|\Z)', re.S)
        match = pattern.search(source)
        if match is None:
            report.add("错误", "08_Tier档位", line, "tier", f"现有 Tier{tier} Resource 不存在")
            continue
        body = match.group(2)
        for from_field, to_field in TIER_TO_GODOT.items():
            raw = row.get(from_field, "")
            if not raw and from_field == "pk_down_threshold" and tier == "0":
                raw = "0"
            if not numeric(raw, "float"):
                continue
            val = float(raw) * (PK_PERCENT_POINT_SCALE if from_field.startswith("pk_") else 1.0)
            replacement = str(int(val)) if from_field == "repeat_count" else format(val, ".15g")
            expr = re.compile(r"(?m)^" + re.escape(to_field) + r" = [^\r\n]*")
            if expr.search(body) is None:
                if from_field in MOTION_WEIGHT_FIELDS:
                    # 现有Tier资源省略了运动字段，Godot脚本仍有export属性，预览显式写出权重。
                    body = body.rstrip() + "\n" + to_field + " = " + replacement + "\n"
                    continue
                report.add("错误", "08_Tier档位", line, from_field, f"正式模板缺字段 {to_field}")
                continue
            body = expr.sub(to_field + " = " + replacement, body, count=1)
        source = source[:match.start(2)] + body + source[match.end(2):]
    return source


def make_combat_numerics_resource(records: list, report: Report) -> str:
    """将 09 表的正式上涨值写入战斗配置，避免胜利事件继续读试玩占位配置。"""
    source = (ROOT / "data/combat_stage/combat_numerics.tres").read_text(encoding="utf-8")
    rules = {row.get("rule_id", ""): (line, row) for line, row in records}
    mapping = {
        "contradiction_break": (
            "contradiction_break_viewer_gain", "contradiction_break_like_gain",
            "contradiction_break_boost_duration_seconds",
        ),
        "oracle_confirm": (
            "oracle_confirmation_viewer_gain", "oracle_confirmation_like_gain",
            "oracle_confirmation_boost_duration_seconds",
        ),
    }
    fields = {
        "contradiction_break_viewer_gain": "viewer_value",
        "contradiction_break_like_gain": "like_value",
        "contradiction_break_boost_duration_seconds": "duration_s",
        "oracle_confirmation_viewer_gain": "viewer_value",
        "oracle_confirmation_like_gain": "like_value",
        "oracle_confirmation_boost_duration_seconds": "duration_s",
    }
    for rule_id, target_names in mapping.items():
        entry = rules.get(rule_id)
        if entry is None:
            report.add("错误", "09_直播数据", 1, "rule_id", f"缺少已配置的正式规则 {rule_id}")
            continue
        line, row = entry
        for target_name in target_names:
            source_name = fields[target_name]
            raw = row.get(source_name, "")
            if not numeric(raw, "float"):
                report.add("错误", "09_直播数据", line, source_name, "正式上涨配置必须为数值")
                continue
            number = float(raw)
            if target_name.endswith(("viewer_gain", "like_gain")) and not number.is_integer():
                report.add("错误", "09_直播数据", line, source_name, "直播上涨人数必须为整数")
                continue
            value = str(int(number)) if number.is_integer() else format(number, ".15g")
            pattern = re.compile(r"(?m)^" + re.escape(target_name) + r" = [^\r\n]*")
            replacement = target_name + " = " + value
            if pattern.search(source) is None:
                source = source.replace("[resource]", "[resource]\n" + replacement, 1)
            else:
                source = pattern.sub(replacement, source, count=1)
    return source



def make_test_only_levels(tables: dict, report: Report) -> dict[str, str]:
    """将 02/03/04/05 的测试记录组合成现有 LevelProfile + LevelCatalog。

    只在 --test-only 下调用；自动使用 test_ 稳定 ID，不把临时参数写入正式关卡。
    """
    root = "data/test_only/generated/level_configuration"
    levels = [r for _, r in tables["02_主播关卡"][1] if r.get("enabled", "true") != "false"]
    generation = {r["level_id"]: r for _, r in tables["05_关卡生成"][1]}
    contradictions = [r for _, r in tables["04_矛盾内容"][1] if r.get("enabled", "true") != "false"]
    pools = {r["pool_id"] for _, r in tables["03_普通词库"][1] if r.get("enabled", "true") == "true"}
    output: dict[str, str] = {}
    order_list: list[tuple[int, str]] = []
    for row_number, row in tables["02_主播关卡"][1]:
        if row not in levels:
            continue
        level_id = row["level_id"]
        streamer_id = row["streamer_id"]
        pool_id = row.get("word_pool_id", "")
        set_id = row.get("contradiction_set_id", "")
        # TEST_ONLY 自建的稳定 ID 带 test_；避免真实业务配置混入这套样例。
        if not all(v.startswith("test_") for v in (level_id, streamer_id, pool_id, set_id)):
            report.add("错误", "02_主播关卡", row_number, "test_only", "测试关卡/主播/词库/矛盾组 ID 必须以 test_ 开头")
            continue
        if pool_id not in pools or level_id not in generation:
            report.add("错误", "02_主播关卡", row_number, "word_pool_id", "缺少同 ID 的词库或关卡生成配置")
            continue
        related = [c for c in contradictions if c.get("set_id") == set_id and c.get("streamer_id") == streamer_id]
        if {c.get("type") for c in related} != {"true", "false"}:
            report.add("错误", "04_矛盾内容", row_number, "set_id", f"{set_id} 缺少真假矛盾配对")
            continue
        if not numeric(row["level_order"], "int") or not re.fullmatch(r"test_[A-Za-z0-9_]+", level_id):
            report.add("错误", "02_主播关卡", row_number, "level_id", "测试关卡 ID 或顺序非法")
            continue
        gen = generation[level_id]
        script = [
            '[gd_resource type="Resource" script_class="LevelProfile" format=3]',
            "",
            '[ext_resource type="Script" path="res://data/level_configuration/level_profile.gd" id="1_profile"]',
            '[ext_resource type="Script" path="res://data/level_configuration/level_contradiction.gd" id="2_contradiction"]',
            f'[ext_resource type="Resource" path="res://{root}/{pool_id}.tres" id="3_pool"]',
            "",
        ]
        true_refs = []
        false_refs = []
        clues = []
        for i, contradiction in enumerate(related, 1):
            resource_id = f"Contradiction_{i}"
            script.extend([
                f'[sub_resource type="Resource" script_class="LevelContradiction" id="{resource_id}"]',
                'script = ExtResource("2_contradiction")',
                'original_sentence_id = ' + gd_string(contradiction["contradiction_id"]),
                'text = ' + gd_string(contradiction["text"]),
                "",
            ])
            ref = f'SubResource("{resource_id}")'
            (true_refs if contradiction["type"] == "true" else false_refs).append(ref)
            if contradiction.get("clue_text", "").strip():
                clues.append(contradiction["clue_text"])
        script += [
            "[resource]",
            f'resource_name = {gd_string("TEST_ONLY " + level_id)}',
            'script = ExtResource("1_profile")',
            f'level_id = {gd_string(level_id)}',
            f'level_order = {int(float(row["level_order"]))}',
            f'streamer_id = {gd_string(streamer_id)}',
            f'streamer_name = {gd_string(row.get("streamer_name", ""))}',
            f'stream_topic = {gd_string(row.get("live_theme", ""))}',
            'normal_speech_pool_source = ExtResource("3_pool")',
            # TEST_ONLY 临时混合权重，仅用于让四种话语都能被抽到。
            "orthodox_ratio = 1.0",
            "heretical_ratio = 1.0",
            "absurd_ratio = 1.0",
            "neutral_ratio = 0.5",
            'true_contradictions = Array[ExtResource("2_contradiction")]([' + ", ".join(true_refs) + "])",
            'false_contradictions = Array[ExtResource("2_contradiction")]([' + ", ".join(false_refs) + "])",
            'contradiction_context_clues = Array[String]([' + ", ".join(gd_string(x) for x in clues) + "])",
        ]
        for from_field, to_field, kind in (
            ("base_batch_count", "base_batch_count", "int"),
            ("spawn_interval_s", "base_spawn_interval_seconds", "float"),
            ("move_speed_px_s", "base_move_speed_pixels_per_second", "float"),
            ("normal_screen_cap", "normal_barrage_screen_cap", "int"),
        ):
            raw = gen.get(from_field, "")
            if not numeric(raw, kind):
                report.add("错误", "05_关卡生成", row_number, from_field, "TEST_ONLY 关卡需填写有效数值")
                continue
            value = str(int(float(raw))) if kind == "int" else format(float(raw), ".15g")
            script.append(f"{to_field} = {value}")
        output[f"{root}/{level_id}.tres"] = "\n".join(script) + "\n"
        order_list.append((int(float(row["level_order"])), level_id))
    if len(order_list) != len(levels) or len({x[0] for x in order_list}) != len(order_list):
        report.add("错误", "02_主播关卡", 1, "level_order", "TEST_ONLY 关卡未能完整生成或顺序重复")
    ordered = [x[1] for x in sorted(order_list)]
    # 这里创建真实 LevelCatalog，后续由 Sandbox 的导出属性选择加载。
    catalog = [
        '[gd_resource type="Resource" script_class="LevelCatalog" format=3]', "",
        '[ext_resource type="Script" path="res://data/level_configuration/level_catalog.gd" id="1_catalog"]',
    ]
    for i, name in enumerate(ordered, 2):
        catalog.append(f'[ext_resource type="Resource" path="res://{root}/{name}.tres" id="{i}_level"]')
    catalog += [
        "", "[resource]", 'resource_name = "TEST_ONLY two-level integration catalog"',
        'script = ExtResource("1_catalog")',
        'profiles = Array[ExtResource("2_level")]([' + ", ".join(f'ExtResource("{i}_level")' for i in range(2, len(ordered) + 2)) + "])",
        "",
    ]
    output[f"{root}/test_only_level_catalog.tres"] = "\n".join(catalog)
    return output


def read_approved_csv_tables(source_root: Path, report: Report) -> dict:
    """读取已批准的正式 CSV；按原始表头和行序进入现有校验链。"""
    tables = {}
    for sheet in SHEET_NAMES:
        path = source_root / f"{sheet}.csv"
        if not path.is_file():
            report.add("错误", sheet, 1, "CSV", f"缺少正式来源文件：{path}")
            continue
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            header = reader.fieldnames or []
            required = RULES[sheet][1]
            if not is_header(header, required):
                report.add("错误", sheet, 1, "表头", "CSV 缺少正式字段或表头无效")
                continue
            records = [(line, dict(row)) for line, row in enumerate(reader, 2)]
        tables[sheet] = (header, records)
        report.stats.append((sheet, len(records), 0, 0))
    validate(tables, report)
    return tables


def make_formal_levels(tables: dict, report: Report) -> dict[str, str]:
    """把正式 02/03/04/05/11 表记录组合为生产关卡 Resource 与目录。"""
    root = "data/level_configuration"
    levels = [(line, row) for line, row in tables[SHEET_NAMES[1]][1]
              if row.get("enabled", "true").lower() == "true"]
    generation = {row["level_id"]: row for _, row in tables[SHEET_NAMES[4]][1]}
    paradox = {row["level_id"]: row for _, row in tables[SHEET_NAMES[12]][1]}
    global_combat_params = {row["param_id"]: row.get("value", "") for _, row in tables[SHEET_NAMES[5]][1]}
    speeches = [row for _, row in tables[SHEET_NAMES[2]][1]
                if row.get("enabled", "").lower() == "true"]
    contradictions = [row for _, row in tables[SHEET_NAMES[3]][1]
                      if row.get("enabled", "").lower() == "true"]
    portraits = {row["portrait_set_id"]: row for _, row in tables[SHEET_NAMES[24]][1]
                 if row.get("enabled", "").lower() == "true"}
    output: dict[str, str] = {}
    order_list = []
    for line, row in levels:
        level_id, streamer_id = row["level_id"], row["streamer_id"]
        pool_id, set_id = row.get("word_pool_id", ""), row.get("contradiction_set_id", "")
        gen, paradox_row = generation.get(level_id), paradox.get(level_id)
        if gen is None or paradox_row is None:
            report.add("错误", "02_主播关卡", line, "level_id", "缺少同 ID 的 05/11 关卡参数")
            continue
        pool_rows = [speech for speech in speeches if speech.get("pool_id") == pool_id]
        related = [c for c in contradictions if c.get("set_id") == set_id and c.get("streamer_id") == streamer_id]
        if not pool_rows:
            report.add("错误", "02_主播关卡", line, "word_pool_id", f"词库 {pool_id!r} 没有启用句子")
            continue
        if {c.get("type") for c in related} != {"true", "false"}:
            report.add("错误", "04_矛盾内容", line, "set_id", f"{set_id} 缺少真假矛盾配对")
            continue
        if not numeric(row.get("level_order", ""), "int"):
            report.add("错误", "02_主播关卡", line, "level_order", "关卡顺序必须为整数")
            continue
        script = [
            '[gd_resource type="Resource" script_class="LevelProfile" format=3]', "",
            '[ext_resource type="Script" path="res://data/level_configuration/level_profile.gd" id="1_profile"]',
            '[ext_resource type="Script" path="res://data/level_configuration/level_contradiction.gd" id="2_contradiction"]',
            f'[ext_resource type="Resource" path="res://{root}/pools/{pool_id}.tres" id="3_pool"]',
        ]
        portrait = portraits.get(row.get("portrait_set_id", ""))
        portrait_resources = {}
        for source_field, profile_field, resource_id in (
            ("idle_asset_path", "streamer_portrait", "4_portrait"),
            ("tier_01_asset_path", "opponent_tier_01_portrait", "5_portrait_t1"),
            ("tier_02_asset_path", "opponent_tier_02_portrait", "6_portrait_t2"),
            ("tier_03_asset_path", "opponent_tier_03_portrait", "7_portrait_t3"),
            ("defeat_asset_path", "opponent_defeat_portrait", "8_portrait_defeat"),
            ("defeat_transition_asset_path", "opponent_defeat_transition_portrait", "9_portrait_transition"),
        ):
            asset_path = portrait.get(source_field, "") if portrait else ""
            if asset_path and (ROOT / asset_path).is_file():
                script.append(f'[ext_resource type="Texture2D" path="res://{asset_path}" id="{resource_id}"]')
                portrait_resources[profile_field] = resource_id
        script.append("")
        true_refs, false_refs, clues = [], [], []
        for i, contradiction in enumerate(related, 1):
            rid = f"Contradiction_{i}"
            script.extend([
                f'[sub_resource type="Resource" script_class="LevelContradiction" id="{rid}"]',
                'script = ExtResource("2_contradiction")',
                f'original_sentence_id = {gd_string(contradiction["contradiction_id"])}',
                f'text = {gd_string(contradiction["text"])}', "",
            ])
            (true_refs if contradiction["type"] == "true" else false_refs).append(f'SubResource("{rid}")')
            clue = contradiction.get("clue_text", "").strip()
            if clue:
                clues.append(clue)
        script.extend([
            "[resource]", 'script = ExtResource("1_profile")',
            f'level_id = {gd_string(level_id)}',
            f'level_order = {int(float(row["level_order"]))}',
            f'streamer_id = {gd_string(streamer_id)}',
            f'streamer_name = {gd_string(row.get("streamer_name", ""))}',
            f'fan_group_name = {gd_string(row.get("fan_group_name", ""))}',
            f'stream_topic = {gd_string(row.get("live_theme", ""))}',
            f'portrait_set_id = {gd_string(row.get("portrait_set_id", ""))}',
            f'tutorial_config_id = {gd_string(row.get("tutorial_config_id", ""))}',
            f'story_config_id = {gd_string(row.get("story_config_id", ""))}',
            f'contradiction_set_id = {gd_string(set_id)}',
            f'background_asset_id = {gd_string(row.get("background_asset_id", ""))}',
            f'fan_badge_asset_id = {gd_string(row.get("fan_badge_asset_id", ""))}',
            'normal_speech_pool_source = ExtResource("3_pool")',
        ])
        for profile_field, resource_id in portrait_resources.items():
            script.append(f'{profile_field} = ExtResource("{resource_id}")')
        # 表中未配置类别比例；用该关卡实际词池的分类权重总和作为类别抽取权重。
        category_weights = defaultdict(float)
        for speech in pool_rows:
            tendency = "heretical" if speech["tendency"] == "heresy" else speech["tendency"]
            category_weights[tendency] += float(speech["weight"])
        for tendency in ("orthodox", "heretical", "absurd", "neutral"):
            script.append(f'{tendency}_ratio = {format(category_weights[tendency], ".15g")}')
        script.extend([
            f'true_contradictions = Array[ExtResource("2_contradiction")]([{", ".join(true_refs)}])',
            f'false_contradictions = Array[ExtResource("2_contradiction")]([{", ".join(false_refs)}])',
            f'contradiction_context_clues = Array[String]([{", ".join(gd_string(x) for x in clues)}])',
        ])
        for source_field, target_field, kind in (
            ("base_batch_count", "base_batch_count", "int"),
            ("spawn_interval_s", "base_spawn_interval_seconds", "float"),
            ("move_speed_px_s", "base_move_speed_pixels_per_second", "float"),
            ("word_weight_multiplier", "word_weight_multiplier", "float"),
        ):
            raw = gen.get(source_field, "")
            if not numeric(raw, kind):
                report.add("错误", "05_关卡生成", line, source_field, "必须有有效的正式数值")
                continue
            value = str(int(float(raw))) if kind == "int" else format(float(raw), ".15g")
            script.append(f"{target_field} = {value}")
        for source_field in ("valid_area_key", "layout_key"):
            if gen.get(source_field, "").strip():
                script.append(f"{source_field} = {gd_string(gen[source_field])}")
        # 05 可覆盖单关特殊实例上限；留空时使用 06 的正式全局值。
        cap = gen.get("special_instance_screen_cap", "").strip()
        if not cap:
            cap = global_combat_params.get("special_foreground_instance_cap", "").strip()
        if cap and numeric(cap, "int"):
            script.append(f"special_instance_screen_cap = {int(float(cap))}")
        else:
            report.add("警告", "05_关卡生成", line, "special_instance_screen_cap",
                       "单关覆盖和 06 全局默认值均为空；保留 LevelProfile 默认值 -1")
        for source_field, target_field, kind in (
            ("time_limit_s", "paradox_window_seconds", "float"),
            ("shot_limit", "paradox_shot_limit", "int"),
            ("silence_transition_s", "paradox_silence_transition_seconds", "float"),
            ("candidate_count", "paradox_candidate_count", "int"),
            ("true_count", "paradox_true_count", "int"),
            ("fake_count", "paradox_fake_count", "int"),
            ("motion_duration_s", "paradox_motion_duration_seconds", "float"),
        ):
            raw = paradox_row.get(source_field, "").strip()
            if raw and numeric(raw, kind):
                value = str(int(float(raw))) if kind == "int" else format(float(raw), ".15g")
                script.append(f"{target_field} = {value}")
        for source_field, target_field in (
            ("motion_mode", "paradox_motion_mode"),
            ("motion_speed_curve_id", "paradox_motion_speed_curve_id"),
        ):
            if paradox_row.get(source_field, "").strip():
                script.append(f"{target_field} = {gd_string(paradox_row[source_field])}")
        shootable = paradox_row.get("shootable_during_motion", "").strip()
        if shootable:
            script.append(f"paradox_shootable_during_motion = {shootable.lower()}")
        output[f"{root}/{level_id}.tres"] = "\n".join(x for x in script if x) + "\n"
        order_list.append((int(float(row["level_order"])), level_id))
    if len(order_list) != len(levels) or len({order for order, _ in order_list}) != len(order_list):
        report.add("错误", "02_主播关卡", 1, "level_order", "四关资源未完整生成或顺序重复")
    ordered = [level_id for _, level_id in sorted(order_list)]
    catalog = [
        '[gd_resource type="Resource" script_class="LevelCatalog" format=3]', "",
        '[ext_resource type="Script" path="res://data/level_configuration/level_catalog.gd" id="1_catalog"]',
    ]
    for index, level_id in enumerate(ordered, 2):
        catalog.append(f'[ext_resource type="Resource" path="res://{root}/{level_id}.tres" id="{index}_level"]')
    catalog.extend(["", "[resource]", 'script = ExtResource("1_catalog")',
        'profiles = Array[ExtResource("2_level")]( [' + ", ".join(
            f'ExtResource("{index}_level")' for index in range(2, len(ordered) + 2)) + "])", ""])
    output[f"{root}/level_catalog.tres"] = "\n".join(catalog)
    return output

def main() -> int:
    parser = argparse.ArgumentParser(description="一次性导出全部策划 Sheet 至分系统 CSV + 已接入的 Godot Resource")
    parser.add_argument("--input", type=Path, help="Google Sheets 下载的 .xlsx")
    parser.add_argument("--approved-csv", action="store_true", help="Generate production Resources from approved data/source_tables CSVs")
    parser.add_argument("--csv-dir", type=Path, help="CSV source directory for --approved-csv")
    parser.add_argument("--output-root", type=Path, help="导出根目录；省略时写入当前项目根目录")
    parser.add_argument("--csv-only", action="store_true", help="仅导出 CSV，不更新已接入的 Resource")
    parser.add_argument("--test-only", action="store_true", help="输入测试工作簿；仅输出 TEST_ONLY CSV/词库，保留正式数据")
    args = parser.parse_args()
    output_root = args.output_root.expanduser().resolve() if args.output_root else ROOT
    report = Report()
    if args.approved_csv:
        if args.test_only:
            print("--approved-csv cannot be combined with --test-only", file=sys.stderr)
            return 2
        source_root = args.csv_dir.expanduser().resolve() if args.csv_dir else ROOT / "data/source_tables"
        tables = read_approved_csv_tables(source_root, report)
        outputs: dict[str, str] = {}
        if not report.errors:
            outputs.update(make_speech_pools(tables[SHEET_NAMES[2]][1], report, "data/level_configuration/pools"))
            outputs.update(make_formal_levels(tables, report))
            outputs["data/combat_stage/tier_catalog.tres"] = make_tier_resource(tables[SHEET_NAMES[9]][1], report)
            outputs["data/combat_stage/combat_numerics.tres"] = make_combat_numerics_resource(tables[SHEET_NAMES[10]][1], report)
        for sheet, n, blanks, examples in report.stats:
            print(f"[table] {sheet}: {n} rows")
        for warning in report.warnings:
            print("[warning]", warning)
        for error in report.errors:
            print("[error]", error, file=sys.stderr)
        if report.errors:
            print(f"Formal CSV import failed: {len(report.errors)} errors; no generated files written.", file=sys.stderr)
            return 1
        for relative, content in outputs.items():
            dest = output_root / relative
            dest.parent.mkdir(parents=True, exist_ok=True)
            if dest.exists() and dest.read_text(encoding="utf-8") == content:
                print("[unchanged]", relative)
            else:
                dest.write_text(content, encoding="utf-8", newline="\n")
                print("[wrote]", relative)
        print(f"Formal CSV import complete: {len(tables)} tables, {len(outputs)} outputs, {len(report.warnings)} warnings, 0 errors.")
        return 0
    if args.input is None:
        parser.error("XLSX mode requires --input; use --approved-csv for approved CSV data")
    source = args.input.expanduser().resolve()
    if not source.is_file():
        print(f"错误：找不到输入 XLSX：{source}", file=sys.stderr)
        return 2
    from openpyxl import load_workbook
    workbook = load_workbook(source, read_only=True, data_only=True)
    # 两种模式共享 A1 来源标记判定，防止测试源和正式来源互相串写。
    has_test_only_marker = (
        "00_填写说明" in workbook.sheetnames
        and "TEST_ONLY" in str(workbook["00_填写说明"]["A1"].value or "")
    )
    if args.test_only and not has_test_only_marker:
        print("错误：--test-only 只接受带 TEST_ONLY 声明的测试工作簿。", file=sys.stderr)
        workbook.close()
        return 2
    if not args.test_only and has_test_only_marker:
        print("错误：普通导表模式拒绝带 TEST_ONLY 声明的工作簿；请使用 --test-only。", file=sys.stderr)
        workbook.close()
        return 2
    tables: dict[str, tuple[list[str], list]] = {}
    # 新版 Google Sheet 将原单页 06 的三种不同结构拆成独立工作表。
    # 旧 TEST_ONLY 夹具仍使用单页区块格式，保留兼容读取。
    split06_parts = ("06_生命周期", "06_基础参数")
    has_any_split06 = any(name in workbook.sheetnames for name in split06_parts)
    split06 = all(name in workbook.sheetnames for name in split06_parts)
    if has_any_split06 and not split06:
        report.add("错误", "06_战斗数值", 1, "Sheet", "拆分 06 工作表不完整：须同时包含生命周期和基础参数")
    for sheet in SHEET_NAMES:
        if sheet in split06_parts and not split06:
            continue
        if sheet not in workbook.sheetnames:
            if sheet in ADDITIONAL_SHEETS:
                continue  # 新增19～25可选，允许旧TEST_ONLY与历史正式工作簿继续导表。
            report.add("错误", sheet, 1, "Sheet", "缺少工作表")
            continue
        # 只有采用新版22表的正式工作簿启用原生类型审计，不破坏历史TEST_ONLY夹具。
        if not args.test_only and "22_直播评论表现参数" in workbook.sheetnames:
            validate_native_boolean_cells(workbook[sheet], report)
        rows = xlsx_rows(workbook[sheet])
        if sheet == "06_战斗数值" and not split06:
            # 三个独立字段区块，标题与第二/第三个表头均不能作为数据输出。
            cut1 = next((i for i, (_, r) in enumerate(rows) if r and r[0].strip() == "【生命周期】"), -1)
            cut2 = next((i for i, (_, r) in enumerate(rows) if r and r[0].strip() == "【系统案补充基础参数】"), -1)
            if not (0 < cut1 < cut2):
                report.add("错误", sheet, 1, "区块", "缺少或排列错误：生命周期、系统案补充基础参数")
                continue
            for name, start, end in (
                ("06_战斗数值", 0, cut1), ("06_生命周期", cut1 + 1, cut2), ("06_基础参数", cut2 + 1, None),
            ):
                tables[name] = parse_sheet(name, rows, report, start, end)
        else:
            tables[sheet] = parse_sheet(sheet, rows, report)
    workbook.close()
    validate(tables, report)
    outputs: dict[str, str] = {}
    for name, (header, records) in tables.items():
        if not header:
            continue
        if args.test_only:
            # 01 和 06/08 已有确定来源：测试工作簿只复制用于引用校验，避免产生第二份权威数据。
            if name.startswith(("01_", "06_", "08_")):
                continue
            outputs[f"data/test_only/source_tables/{name}.csv"] = csv_body(header, records)
        else:
            outputs[f"data/source_tables/{name}.csv"] = csv_body(header, records)
    if not args.csv_only and "03_普通词库" in tables and not report.errors:
        if args.test_only:
            speech_root = "data/test_only/generated/level_configuration"
        else:
            # 普通词库现阶段尚未批准，原表预览不能进入正式 Resource 目录。
            speech_root = "data/test_only/sheet_preview/level_configuration"
        outputs.update(make_speech_pools(tables[SHEET_NAMES[2]][1], report, speech_root))
        if args.test_only:
            outputs.update(make_test_only_levels(tables, report))
        if not args.test_only and "08_Tier档位" in tables:
            outputs["data/test_only/sheet_preview/combat_stage/tier_catalog.tres"] = make_tier_resource(tables[SHEET_NAMES[9]][1], report)
    for sheet, n, blanks, examples in report.stats:
        print(f"[表] {sheet}：{n} 条，空行 {blanks}，跳过说明/示例 {examples}")
    for warning in report.warnings:
        print("[警告]", warning)
    for error in report.errors:
        print("[错误]", error, file=sys.stderr)
    if report.errors:
        print(f"导表失败：{len(report.errors)} 个错误；未写入生成文件。", file=sys.stderr)
        return 1
    # 校验全部通过后才覆盖工具拥有的文件；同名之外的仓库资源绝不修改。
    for relative, content in outputs.items():
        dest = output_root / relative
        dest.parent.mkdir(parents=True, exist_ok=True)
        if dest.exists() and dest.read_text(encoding="utf-8") == content:
            print("[未变化]", relative)
        else:
            dest.write_text(content, encoding="utf-8", newline="\n")
            print("[写入]", relative)
    print(f"导表完成：{len(tables)} 张结构表，{len(outputs)} 个文件，"
          f"{len(report.warnings)} 条警告，0 个错误。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
