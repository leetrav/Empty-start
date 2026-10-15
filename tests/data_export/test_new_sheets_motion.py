"""新增业务Sheet、混合配置、Tier运动权重与预览分档的独立回归。"""
from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from tools import export_game_data as tool


def sample(sheet: str, records: list[dict]) -> tuple:
    header = list(tool.RULES[sheet][1])
    return header, [(n + 2, record) for n, record in enumerate(records)]


def main() -> None:
    extras = tuple(tool.SHEET_NAMES[-7:])
    assert extras == (
        "19_直播评论词库", "20_主播气泡对白", "21_反击技能池",
        "22_直播评论表现参数", "23_对手立绘配置",
        "24_新手教学配置", "25_直播观众账号",
    )

    # 一列允许浮点时长和字符串true/false；非法字符串仍是硬错误。
    mixed = ("22_直播评论表现参数", (
        ["param_id", "value"],
        [(2, {"param_id": "time", "value": "2.5"}),
         (3, {"param_id": "comment_pause_with_game", "value": "true"}),
         (4, {"param_id": "comment_updates_business_counter", "value": "false"})],
    ))
    report = tool.Report()
    tool.validate({mixed[0]: mixed[1]}, report)
    assert not report.errors, report.errors
    bad = tool.Report()
    tool.validate({mixed[0]: (
        mixed[1][0], mixed[1][1] + [(5, {"param_id": "bad", "value": "random"})]
    )}, bad)
    assert bad.errors and "value" in bad.errors[0]

    # 迁移错误回归：源xlsx文本true不是原生布尔值，正式工作簿应阻断。
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "22_直播评论表现参数"
    ws.append(["param_id", "value", "unit"])
    ws.append(["comment_pause_with_game", "true", "bool"])
    ws.append(["comment_updates_business_counter", False, "bool"])
    wrong_type = tool.Report()
    tool.validate_native_boolean_cells(ws, wrong_type)
    assert len(wrong_type.errors) == 1, wrong_type.errors
    ws["B2"] = True
    valid_type = tool.Report()
    tool.validate_native_boolean_cells(ws, valid_type)
    assert not valid_type.errors, valid_type.errors
    ws.append(["comment_visible_row_cap", True, "rows"])
    invalid_numeric = tool.Report()
    tool.validate_native_boolean_cells(ws, invalid_numeric)
    assert len(invalid_numeric.errors) == 1, invalid_numeric.errors
    ws["B4"] = 5
    ws.title = "04_矛盾内容"
    ws.delete_rows(1, ws.max_row)
    ws.append(["contradiction_id", "enabled"])
    ws.append(["c1", "true"])
    invalid_enabled = tool.Report()
    tool.validate_native_boolean_cells(ws, invalid_enabled)
    assert len(invalid_enabled.errors) == 1, invalid_enabled.errors
    ws["B2"] = True
    valid_enabled = tool.Report()
    tool.validate_native_boolean_cells(ws, valid_enabled)
    assert not valid_enabled.errors, valid_enabled.errors
    wb.close()

    # 新表主键和关联验证；停用的历史 streamer_a 样例不阻断正式导表。
    tables = {
        "02_主播关卡": (
            ["level_id", "level_order", "streamer_id", "portrait_set_id",
             "tutorial_config_id", "story_config_id"],
            [(2, {"level_id": "level_003", "level_order": "3",
                  "streamer_id": "frog", "portrait_set_id": "portrait_frog",
                  "tutorial_config_id": "tutorial_one", "story_config_id": "story_frog"})],
        ),
        "23_对手立绘配置": (
            ["portrait_set_id", "streamer_id", "idle_asset_path", "enabled"],
            [(2, {"portrait_set_id": "portrait_frog", "streamer_id": "frog",
                  "idle_asset_path": "assets/characters/opponents/frog/frog_idle.png",
                  "enabled": "true"})],
        ),
        "24_新手教学配置": (
            ["tutorial_config_id", "step_id", "step_order", "enabled"],
            [(2, {"tutorial_config_id": "tutorial_one", "step_id": "step_one",
                  "step_order": "1", "enabled": "true"})],
        ),
        "20_主播气泡对白": (
            ["dialogue_id", "streamer_id", "story_config_id",
             "speaker_side", "trigger_type", "text", "enabled"],
            [(2, {"dialogue_id": "line_frog", "streamer_id": "frog",
                  "story_config_id": "story_frog", "speaker_side": "opponent",
                  "trigger_type": "connect", "text": "hello", "enabled": "true"}),
             (3, {"dialogue_id": "old", "streamer_id": "streamer_a",
                  "story_config_id": "", "speaker_side": "opponent",
                  "trigger_type": "connect", "text": "disabled example", "enabled": "false"})],
        ),
    }
    valid = tool.Report()
    tool.validate(tables, valid)
    assert not valid.errors, valid.errors
    wrong = {**tables, "02_主播关卡": (
        tables["02_主播关卡"][0],
        [(2, {**tables["02_主播关卡"][1][0][1],
              "portrait_set_id": "not_real"})],
    )}
    badlink = tool.Report()
    tool.validate(wrong, badlink)
    assert any("not_real" in s for s in badlink.errors), badlink.errors

    # 运行旧Resource时只有Tier独立区块才可覆写；四个运动权重必须分档。
    motion = [(70, 20, 5, 5), (60, 25, 10, 5),
              (50, 25, 15, 10), (45, 30, 15, 10),
              (40, 30, 20, 10), (35, 30, 25, 10)]
    records = []
    for tier, weights in enumerate(motion):
        data = {
            "tier": str(tier), "pk_up_threshold": str([55, 63, 70, 78, 89, 100][tier]),
            "pk_down_threshold": "" if tier == 0 else str([0, 55, 63, 70, 78, 89][tier]),
            "spawn_batch_mult": "1", "spawn_frequency_mult": "1",
            "move_speed_mult": "1", "lifetime_mult": "1",
            "pullback_mult": "0" if tier == 0 else str(tier / 5),
            "repeat_count": "1", "repeat_lifetime_s": "3",
        }
        data.update(dict(zip(tool.MOTION_WEIGHT_FIELDS, map(str, weights))))
        records.append((tier + 2, data))
    report = tool.Report()
    tool.validate({"08_Tier档位": (
        list(records[0][1]), records
    )}, report)
    assert not report.errors, report.errors
    generated = tool.make_tier_resource(records, report)
    assert not report.errors, report.errors
    parts = re.findall(
        r'(?ms)^\[sub_resource type="Resource" id="Tier(\d)"\]\n'
        r'(.*?)(?=^\[(?:sub_resource|resource)\b|\Z)',
        generated,
    )
    assert len(parts) == 6, len(parts)
    for n, (label, body) in enumerate(parts):
        assert label == str(n)
        for key, weight in zip(
            ("straight_motion_weight", "curved_motion_weight",
             "speed_change_motion_weight", "wandering_motion_weight"),
            motion[n],
        ):
            assert re.search(r"(?m)^" + key + r" = " + str(weight) + r"$", body), (label, key)
        assert body.count("straight_motion_weight =") == 1
        expected_pullback = str(int(n / 5)) if (n / 5).is_integer() else str(n / 5)
        assert f"opponent_pullback_multiplier = {expected_pullback}" in body, (label, expected_pullback)

    invalid = dict(records[0][1])
    invalid["motion_straight_weight"] = "71"
    bad_tier = tool.Report()
    tool.validate({"08_Tier档位": (
        list(invalid), [(2, invalid)] + records[1:]
    )}, bad_tier)
    assert any("100" in s for s in bad_tier.errors), bad_tier.errors
    print("PASS new sheets, mixed bool/float, cross refs, six-tier motion & isolation")


if __name__ == "__main__":
    main()
